"""Rotas — Motoristas"""
import io
import re as _re
import uuid
import logging
from datetime import date as _date, datetime, timezone
from flask import Blueprint, render_template, request, redirect, url_for, flash, jsonify, send_file
from flask_login import login_required, current_user
from app import db
from app.models.usuario import Motorista, Veiculo, Checklist, MotoristaDocumento
from app.services.sharepoint_upload import (
    upload_documento_motorista, deletar_foto, validar_arquivo,
)

logger      = logging.getLogger(__name__)
bp          = Blueprint('motoristas', __name__, url_prefix='/motoristas')
TAMANHO_MAX = 10 * 1024 * 1024   # 10 MB


def _parse_date(value):
    """Converte string 'YYYY-MM-DD' para date, ou retorna None."""
    if not value:
        return None
    try:
        return _date.fromisoformat(value)
    except (ValueError, TypeError):
        return None


@bp.route('/')
@login_required
def lista():
    q      = request.args.get('q', '')
    ativo  = request.args.get('ativo', '1')
    apto   = request.args.get('apto', '')
    page   = request.args.get('page', 1, type=int)
    query  = Motorista.query
    if q:
        query = query.filter(
            db.or_(Motorista.nome.ilike(f'%{q}%'), Motorista.cpf.ilike(f'%{q}%'))
        )
    if ativo != '':
        query = query.filter_by(ativo=(ativo == '1'))
    if apto != '':
        query = query.filter_by(apto_dirigir=(apto == '1'))
    mots = query.order_by(Motorista.nome).paginate(page=page, per_page=20, error_out=False)
    return render_template('motoristas/lista.html', motoristas=mots, q=q, ativo=ativo, apto=apto)


@bp.route('/<int:id>')
@login_required
def detalhe(id):
    m = Motorista.query.get_or_404(id)
    ultimos_checklists = (
        Checklist.query
        .filter_by(motorista_id=id)
        .order_by(Checklist.data_checklist.desc())
        .limit(5)
        .all()
    )
    return render_template('motoristas/detalhe.html', m=m,
                           ultimos_checklists=ultimos_checklists,
                           tipos_doc=MotoristaDocumento.TIPOS)


@bp.route('/novo', methods=['GET', 'POST'])
@login_required
def novo():
    veiculos = Veiculo.query.filter_by(status='ativo').order_by(Veiculo.placa).all()
    if request.method == 'POST':
        f = request.form
        try:
            m = Motorista(
                nome=f['nome'], cpf=f.get('cpf') or None,
                matricula=f.get('matricula'), telefone=f.get('telefone'),
                email=f.get('email'), municipio=f.get('municipio'),
                cnh_numero=f.get('cnh_numero'), cnh_categoria=f.get('cnh_categoria'),
                cnh_validade=_parse_date(f.get('cnh_validade')),
                cnh_primeira_hab=_parse_date(f.get('cnh_primeira_hab')),
                cnh_pontos=int(f.get('cnh_pontos') or 0),
                veiculo_principal_id=f.get('veiculo_principal_id') or None,
                observacoes=f.get('observacoes'),
            )
            db.session.add(m)
            db.session.commit()
            flash('Motorista cadastrado com sucesso.', 'success')
            return redirect(url_for('motoristas.detalhe', id=m.id))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao cadastrar: {e}', 'danger')
    return render_template('motoristas/form.html', m=None, veiculos=veiculos)


@bp.route('/<int:id>/editar', methods=['GET', 'POST'])
@login_required
def editar(id):
    m = Motorista.query.get_or_404(id)
    veiculos = Veiculo.query.filter(Veiculo.status != 'inativo').order_by(Veiculo.placa).all()
    if request.method == 'POST':
        f = request.form
        try:
            m.nome               = f['nome']
            m.cpf                = f.get('cpf') or None
            m.matricula          = f.get('matricula')
            m.telefone           = f.get('telefone')
            m.email              = f.get('email')
            m.municipio          = f.get('municipio')
            m.cnh_numero         = f.get('cnh_numero')
            m.cnh_categoria      = f.get('cnh_categoria')
            m.cnh_validade       = _parse_date(f.get('cnh_validade'))
            m.cnh_primeira_hab   = _parse_date(f.get('cnh_primeira_hab'))
            m.cnh_pontos         = int(f.get('cnh_pontos') or 0)
            m.veiculo_principal_id = f.get('veiculo_principal_id') or None
            m.ativo              = 'ativo' in f
            m.observacoes        = f.get('observacoes')
            db.session.commit()
            flash('Motorista atualizado.', 'success')
            return redirect(url_for('motoristas.detalhe', id=m.id))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao atualizar: {e}', 'danger')
    return render_template('motoristas/form.html', m=m, veiculos=veiculos)


# ── APTO A DIRIGIR (toggle lateral) ─────────────────────────────────────────
@bp.route('/<int:id>/apto', methods=['POST'])
@login_required
def toggle_apto(id):
    """Alterna se o motorista está apto ou não a dirigir. Usado pelo toggle
    lateral da tela de detalhe — não exige passar pelo formulário completo."""
    m = Motorista.query.get_or_404(id)

    if request.is_json:
        payload = request.get_json(silent=True) or {}
        apto    = bool(payload.get('apto'))
        motivo  = (payload.get('motivo') or '').strip()[:200]
    else:
        apto    = request.form.get('apto') in ('1', 'true', 'True')
        motivo  = (request.form.get('motivo') or '').strip()[:200]

    m.apto_dirigir        = apto
    m.apto_motivo         = motivo or None if not apto else None
    m.apto_atualizado_em  = datetime.now(timezone.utc)
    m.apto_atualizado_por = current_user.id
    db.session.commit()

    if request.is_json:
        return jsonify({
            'ok': True,
            'apto_dirigir': m.apto_dirigir,
            'apto_motivo': m.apto_motivo,
            'atualizado_em': m.apto_atualizado_em.isoformat(),
        })

    flash('Motorista marcado como apto a dirigir.' if apto else
          'Motorista marcado como INAPTO a dirigir.', 'success' if apto else 'warning')
    return redirect(url_for('motoristas.detalhe', id=m.id))


# ── DOCUMENTOS (CNH, termo de direção defensiva, termo de uso de veículo) ───
@bp.route('/<int:id>/documentos', methods=['GET'])
@login_required
def listar_documentos(id):
    Motorista.query.get_or_404(id)
    docs = MotoristaDocumento.query.filter_by(motorista_id=id).all()
    return jsonify([{
        'id':          d.id,
        'tipo':        d.tipo,
        'tipo_nome':   d.tipo_display,
        'nome':        d.nome_arquivo,
        'url':         d.url_sharepoint,
        'created_at':  d.created_at.isoformat() if d.created_at else None,
    } for d in docs])


@bp.route('/<int:id>/documentos/<tipo>', methods=['POST'])
@login_required
def upload_documento(id, tipo):
    """Anexa (ou substitui) o documento de um tipo específico do motorista."""
    m = Motorista.query.get_or_404(id)

    if tipo not in MotoristaDocumento.TIPOS:
        return jsonify({'ok': False, 'erro': 'Tipo de documento inválido.'}), 400

    if 'file' not in request.files:
        return jsonify({'ok': False, 'erro': 'Nenhum arquivo enviado.'}), 400

    arquivo = request.files['file']
    if not arquivo.filename:
        return jsonify({'ok': False, 'erro': 'Nome de arquivo vazio.'}), 400

    ok, erro = validar_arquivo(arquivo.filename, arquivo.mimetype or '')
    if not ok:
        return jsonify({'ok': False, 'erro': erro}), 400

    content = arquivo.read()
    if len(content) > TAMANHO_MAX:
        return jsonify({'ok': False, 'erro': 'Arquivo maior que 10 MB.'}), 400

    filename = f'{uuid.uuid4().hex[:8]}_{arquivo.filename}'[:120]

    result = upload_documento_motorista(
        motorista_nome = m.nome,
        tipo           = tipo,
        filename       = filename,
        content        = content,
        mime_type      = arquivo.mimetype or 'application/octet-stream',
    )
    if not result['ok']:
        logger.error(f'Upload documento motorista {id}/{tipo} falhou: {result["erro"]}')
        return jsonify({'ok': False, 'erro': result['erro']}), 500

    # Substitui documento existente do mesmo tipo (1 anexo ativo por tipo)
    existente = MotoristaDocumento.query.filter_by(motorista_id=id, tipo=tipo).first()
    if existente:
        if existente.sharepoint_id:
            try:
                deletar_foto(existente.sharepoint_id)
            except Exception as e:
                logger.warning(f'Falha ao remover documento antigo do SharePoint: {e}')
        existente.nome_arquivo   = result.get('nome', filename)
        existente.url_sharepoint = result.get('url', '')
        existente.sharepoint_id  = result.get('id', '')
        existente.uploaded_by    = current_user.id
        existente.created_at     = datetime.now(timezone.utc)
        doc = existente
    else:
        doc = MotoristaDocumento(
            motorista_id   = id,
            tipo           = tipo,
            nome_arquivo   = result.get('nome', filename),
            url_sharepoint = result.get('url', ''),
            sharepoint_id  = result.get('id', ''),
            uploaded_by    = current_user.id,
        )
        db.session.add(doc)

    db.session.commit()
    return jsonify({
        'ok':        True,
        'id':        doc.id,
        'tipo':      doc.tipo,
        'tipo_nome': doc.tipo_display,
        'url':       doc.url_sharepoint,
        'nome':      doc.nome_arquivo,
    }), 201


@bp.route('/<int:id>/documentos/<int:doc_id>', methods=['DELETE'])
@login_required
def remover_documento(id, doc_id):
    doc = MotoristaDocumento.query.filter_by(id=doc_id, motorista_id=id).first_or_404()
    if doc.sharepoint_id:
        try:
            deletar_foto(doc.sharepoint_id)
        except Exception as e:
            logger.warning(f'Falha ao remover documento do SharePoint: {e}')
    db.session.delete(doc)
    db.session.commit()
    return jsonify({'ok': True})


# ── EXPORTAR EXCEL ───────────────────────────────────────────────────────────
@bp.route('/exportar')
@login_required
def exportar():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    mots = Motorista.query.order_by(Motorista.nome).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Motoristas'

    hdr_fill  = PatternFill('solid', fgColor='F7931E')
    hdr_font  = Font(bold=True, color='FFFFFF', size=11)
    hdr_align = Alignment(horizontal='center', vertical='center')
    borda     = Border(
        left=Side(style='thin', color='E5E5EA'), right=Side(style='thin', color='E5E5EA'),
        top=Side(style='thin', color='E5E5EA'),  bottom=Side(style='thin', color='E5E5EA'),
    )
    alt_fill = PatternFill('solid', fgColor='F2F2F7')

    COLUNAS = [
        ('nome',              'Nome *',                 30),
        ('cpf',               'CPF',                    18),
        ('matricula',         'Matrícula',               14),
        ('telefone',          'Telefone',                16),
        ('email',             'E-mail',                  28),
        ('municipio',         'Município',               18),
        ('cnh_numero',        'Nº CNH',                  16),
        ('cnh_categoria',     'Categoria CNH',           14),
        ('cnh_validade',      'Validade CNH (AAAA-MM-DD)', 20),
        ('cnh_primeira_hab',  '1ª Habilitação (AAAA-MM-DD)', 22),
        ('cnh_pontos',        'Pontos CNH',              12),
        ('veiculo_placa',     'Placa Veículo Principal', 20),
        ('apto_dirigir',      'Apto a Dirigir (SIM/NAO)', 20),
        ('ativo',             'Ativo (SIM/NAO)',         16),
        ('observacoes',       'Observações',             35),
    ]

    for col_idx, (_, label, width) in enumerate(COLUNAS, start=1):
        cel = ws.cell(row=1, column=col_idx, value=label)
        cel.fill = hdr_fill; cel.font = hdr_font
        cel.alignment = hdr_align; cel.border = borda
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 28

    for row_idx, m in enumerate(mots, start=2):
        valores = [
            m.nome, m.cpf or '', m.matricula or '', m.telefone or '', m.email or '',
            m.municipio or '', m.cnh_numero or '', m.cnh_categoria or '',
            m.cnh_validade.isoformat() if m.cnh_validade else '',
            m.cnh_primeira_hab.isoformat() if m.cnh_primeira_hab else '',
            m.cnh_pontos or 0,
            m.veiculo_principal.placa if m.veiculo_principal else '',
            'SIM' if m.apto_dirigir else 'NAO',
            'SIM' if m.ativo else 'NAO',
            m.observacoes or '',
        ]
        row_fill = alt_fill if row_idx % 2 == 0 else None
        for col_idx, val in enumerate(valores, start=1):
            cel = ws.cell(row=row_idx, column=col_idx, value=val)
            cel.border = borda
            cel.alignment = Alignment(vertical='center')
            if row_fill: cel.fill = row_fill

    ws.freeze_panes = 'A2'
    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)

    return send_file(buf, as_attachment=True,
                     download_name=f'motoristas_{_date.today().isoformat()}.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# ── MODELO (TEMPLATE VAZIO) ──────────────────────────────────────────────────
@bp.route('/modelo')
@login_required
def modelo():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Motoristas'

    hdr_fill = PatternFill('solid', fgColor='F7931E')
    hdr_font = Font(bold=True, color='FFFFFF', size=11)
    borda    = Border(
        left=Side(style='thin', color='E5E5EA'), right=Side(style='thin', color='E5E5EA'),
        top=Side(style='thin', color='E5E5EA'),  bottom=Side(style='thin', color='E5E5EA'),
    )
    COLUNAS = [
        ('Nome *', 30), ('CPF', 18), ('Matrícula', 14), ('Telefone', 16),
        ('E-mail', 28), ('Município', 18), ('Nº CNH', 16), ('Categoria CNH', 14),
        ('Validade CNH (AAAA-MM-DD)', 20), ('1ª Habilitação (AAAA-MM-DD)', 22),
        ('Pontos CNH', 12), ('Placa Veículo Principal', 20),
        ('Apto a Dirigir (SIM/NAO)', 20), ('Ativo (SIM/NAO)', 16), ('Observações', 35),
    ]
    for col_idx, (label, width) in enumerate(COLUNAS, start=1):
        cel = ws.cell(row=1, column=col_idx, value=label)
        cel.fill = hdr_fill; cel.font = hdr_font
        cel.alignment = Alignment(horizontal='center', vertical='center')
        cel.border = borda
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 28

    exemplo_fill = PatternFill('solid', fgColor='FFF3CD')
    for col_idx, val in enumerate(
        ['João da Silva', '000.000.000-00', 'REZ0001', '(93) 9 9999-9999',
         'joao@empresaexemplo.com.br', 'Santarém', '00000000000', 'B',
         '2027-01-31', '2015-06-10', 0, 'ABC1D23', 'SIM', 'SIM', ''], start=1
    ):
        cel = ws.cell(row=2, column=col_idx, value=val)
        cel.fill = exemplo_fill; cel.border = borda
        cel.font = Font(italic=True, color='636366')

    wi = wb.create_sheet('Instruções')
    instrucoes = [
        ('INSTRUÇÕES DE PREENCHIMENTO', True),
        ('', False),
        ('• Nome é obrigatório (marcado com *)', False),
        ('• CPF: usado para identificar duplicatas — se existir, os dados serão ATUALIZADOS', False),
        ('• Sem CPF: motorista sempre criado como novo', False),
        ('• Datas no formato AAAA-MM-DD (ex: 2027-01-31)', False),
        ('• Placa Veículo Principal: deve corresponder a uma placa já cadastrada no sistema', False),
        ('• Apto a Dirigir / Ativo: preencher com SIM ou NAO', False),
        ('• A linha 2 da aba Motoristas é um exemplo — pode apagar antes de importar', False),
        ('• Não altere os cabeçalhos (linha 1)', False),
    ]
    for row_idx, (texto, negrito) in enumerate(instrucoes, start=1):
        cel = wi.cell(row=row_idx, column=1, value=texto)
        cel.font = Font(bold=negrito, size=12 if negrito else 11,
                        color='F7931E' if negrito else '000000')
    wi.column_dimensions['A'].width = 75

    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='modelo_motoristas.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# ── IMPORTAR EXCEL ────────────────────────────────────────────────────────────
@bp.route('/importar', methods=['POST'])
@login_required
def importar():
    """Importa motoristas de uma planilha .xlsx.
    CPF já cadastrado → registro é ATUALIZADO com os dados da linha.
    CPF ausente ou não encontrado → motorista é CRIADO."""
    import openpyxl

    arq = request.files.get('arquivo')
    if not arq or not arq.filename.lower().endswith(('.xlsx', '.xlsm')):
        flash('Envie um arquivo .xlsx válido.', 'danger')
        return redirect(url_for('motoristas.lista'))

    try:
        buf  = io.BytesIO(arq.read())
        wb   = openpyxl.load_workbook(buf, read_only=True, data_only=True)
        ws   = wb['Motoristas'] if 'Motoristas' in wb.sheetnames else wb.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
    except Exception as e:
        flash(f'Erro ao ler o arquivo: {e}', 'danger')
        return redirect(url_for('motoristas.lista'))

    # Cache de placas → id, evita 1 query por linha
    veiculos_por_placa = {
        v.placa.upper(): v.id for v in Veiculo.query.with_entities(Veiculo.id, Veiculo.placa).all()
    }

    def _cpf_normalizado(raw):
        digitos = _re.sub(r'\D', '', str(raw or ''))
        if len(digitos) != 11:
            return str(raw or '').strip() or None
        return f'{digitos[:3]}.{digitos[3:6]}.{digitos[6:9]}-{digitos[9:11]}'

    def _data(raw):
        if not raw:
            return None
        if hasattr(raw, 'date'):
            return raw.date()
        if hasattr(raw, 'isoformat') and not isinstance(raw, str):
            return raw
        try:
            return _date.fromisoformat(str(raw).strip()[:10])
        except (ValueError, TypeError):
            return None

    def _texto(row, idx):
        return str(row[idx] or '').strip() or None if len(row) > idx and row[idx] is not None else None

    criados = atualizados = ignorados = 0
    erros = []

    for linha_num, row in enumerate(rows, start=2):
        if not any(row):
            continue

        nome = str(row[0] or '').strip() if len(row) > 0 else ''
        if not nome:
            erros.append(f'Linha {linha_num}: Nome vazio — ignorada.')
            ignorados += 1
            continue

        cpf = _cpf_normalizado(row[1]) if len(row) > 1 else None

        placa_raw  = _texto(row, 11)
        veiculo_id = veiculos_por_placa.get(placa_raw.upper()) if placa_raw else None
        if placa_raw and not veiculo_id:
            erros.append(f'Linha {linha_num} ({nome[:30]}): placa "{placa_raw}" não encontrada — veículo não vinculado.')

        apto_val  = str(row[12] or 'SIM').strip().upper() != 'NAO' if len(row) > 12 else True
        ativo_val = str(row[13] or 'SIM').strip().upper() != 'NAO' if len(row) > 13 else True

        try:
            cnh_pontos = int(row[10]) if len(row) > 10 and row[10] not in (None, '') else 0
        except (ValueError, TypeError):
            cnh_pontos = 0

        campos = dict(
            nome              = nome,
            matricula         = _texto(row, 2),
            telefone          = _texto(row, 3),
            email             = _texto(row, 4),
            municipio         = _texto(row, 5),
            cnh_numero        = _texto(row, 6),
            cnh_categoria     = _texto(row, 7),
            cnh_validade      = _data(row[8])  if len(row) > 8  else None,
            cnh_primeira_hab  = _data(row[9])  if len(row) > 9  else None,
            cnh_pontos        = cnh_pontos,
            veiculo_principal_id = veiculo_id,
            apto_dirigir      = apto_val,
            ativo             = ativo_val,
            observacoes       = _texto(row, 14),
        )

        try:
            existente = Motorista.query.filter_by(cpf=cpf).first() if cpf else None
            if existente:
                for k, v in campos.items():
                    setattr(existente, k, v)
                atualizados += 1
            else:
                db.session.add(Motorista(cpf=cpf, **campos))
                criados += 1
        except Exception as e:
            db.session.rollback()
            erros.append(f'Linha {linha_num} ({nome[:30]}): {e}')
            ignorados += 1

    try:
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        flash(f'Erro ao salvar: {e}', 'danger')
        return redirect(url_for('motoristas.lista'))

    partes = []
    if criados:     partes.append(f'{criados} criado(s)')
    if atualizados: partes.append(f'{atualizados} atualizado(s)')
    if ignorados:   partes.append(f'{ignorados} ignorado(s)')
    flash(f'Importação concluída: {", ".join(partes)}.', 'success' if not erros else 'warning')
    for msg in erros[:8]:
        flash(msg, 'warning')

    return redirect(url_for('motoristas.lista'))
