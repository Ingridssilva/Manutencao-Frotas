"""Rotas — Administração (usuários, fornecedores, catálogos)"""
from flask import Blueprint, render_template, request, redirect, url_for, flash, abort, send_file
from flask_login import login_required, current_user
from app import db
import io, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from app.models.usuario import (
    Usuario, PerfilAcesso, Fornecedor,
    CatalogoServico, CatalogoMaterial, CatalogoImplemento,
    CategoriaServico, CategoriaMaterial, CategoriaImplemento,
    ChecklistItemConfig
)

bp = Blueprint('admin', __name__, url_prefix='/admin')


def _requer_admin():
    if not current_user.is_authenticated or not current_user.is_admin:
        abort(403)


# ── USUÁRIOS ────────────────────────────────────────────
@bp.route('/usuarios')
@login_required
def usuarios():
    _requer_admin()
    users = Usuario.query.order_by(Usuario.nome).all()
    return render_template('admin/usuarios.html', users=users)


@bp.route('/usuarios/novo', methods=['GET', 'POST'])
@login_required
def usuario_novo():
    _requer_admin()
    perfis = PerfilAcesso.query.all()
    if request.method == 'POST':
        f = request.form
        try:
            u = Usuario(
                nome=f['nome'], email=f['email'].lower().strip(),
                perfil_id=int(f['perfil_id']),
            )
            if f.get('senha'):
                u.set_senha(f['senha'])
            db.session.add(u)
            db.session.commit()
            flash('Usuário criado.', 'success')
            return redirect(url_for('admin.usuarios'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro: {e}', 'danger')
    return render_template('admin/usuario_form.html', u=None, perfis=perfis)


@bp.route('/usuarios/<id>/editar', methods=['GET', 'POST'])
@login_required
def usuario_editar(id):
    _requer_admin()
    u      = Usuario.query.get_or_404(id)
    perfis = PerfilAcesso.query.all()
    if request.method == 'POST':
        f = request.form
        try:
            u.nome      = f['nome']
            u.email     = f['email'].lower().strip()
            u.perfil_id = int(f['perfil_id'])
            u.ativo     = 'ativo' in f
            if f.get('senha'):
                u.set_senha(f['senha'])
            db.session.commit()
            flash('Usuário atualizado.', 'success')
            return redirect(url_for('admin.usuarios'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro: {e}', 'danger')
    return render_template('admin/usuario_form.html', u=u, perfis=perfis)


# ── FORNECEDORES ────────────────────────────────────────
@bp.route('/fornecedores')
@login_required
def fornecedores():
    _requer_admin()
    q     = request.args.get('q', '')
    query = Fornecedor.query
    if q:
        query = query.filter(
            db.or_(Fornecedor.razao_social.ilike(f'%{q}%'),
                   Fornecedor.cnpj.ilike(f'%{q}%'))
        )
    forn = query.order_by(Fornecedor.razao_social).all()
    return render_template('admin/fornecedores.html', fornecedores=forn, q=q)


@bp.route('/fornecedores/novo', methods=['GET', 'POST'])
@bp.route('/fornecedores/<int:id>/editar', methods=['GET', 'POST'])
@login_required
def fornecedor_form(id=None):
    _requer_admin()
    f_obj = Fornecedor.query.get_or_404(id) if id else None
    if request.method == 'POST':
        f = request.form
        try:
            if f_obj:
                f_obj.razao_social  = f['razao_social']
                f_obj.nome_fantasia = f.get('nome_fantasia')
                f_obj.cnpj          = f.get('cnpj') or None
                f_obj.tipo          = f.get('tipo')
                f_obj.telefone      = f.get('telefone')
                f_obj.email         = f.get('email')
                f_obj.endereco      = f.get('endereco')
                f_obj.municipio     = f.get('municipio')
                f_obj.uf            = f.get('uf', 'PA')
                f_obj.observacoes   = f.get('observacoes')
                f_obj.ativo         = 'ativo' in f
            else:
                f_obj = Fornecedor(
                    razao_social=f['razao_social'],
                    nome_fantasia=f.get('nome_fantasia'),
                    cnpj=f.get('cnpj') or None,
                    tipo=f.get('tipo'),
                    telefone=f.get('telefone'),
                    email=f.get('email'),
                    endereco=f.get('endereco'),
                    municipio=f.get('municipio'),
                    uf=f.get('uf', 'PA'),
                    observacoes=f.get('observacoes'),
                )
                db.session.add(f_obj)
            db.session.commit()
            flash('Fornecedor salvo.', 'success')
            return redirect(url_for('admin.fornecedores'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro: {e}', 'danger')
    return render_template('admin/fornecedor_form.html', f=f_obj)



# ── EXPORT FORNECEDORES ──────────────────────────────────
@bp.route('/fornecedores/exportar')
@login_required
def fornecedores_exportar():
    _requer_admin()
    fornecs = Fornecedor.query.order_by(Fornecedor.razao_social).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Fornecedores'

    hdr_fill  = PatternFill('solid', fgColor='F7931E')
    hdr_font  = Font(bold=True, color='FFFFFF', size=11)
    hdr_align = Alignment(horizontal='center', vertical='center')
    borda     = Border(
        left=Side(style='thin', color='E5E5EA'),
        right=Side(style='thin', color='E5E5EA'),
        top=Side(style='thin', color='E5E5EA'),
        bottom=Side(style='thin', color='E5E5EA'),
    )
    alt_fill = PatternFill('solid', fgColor='F2F2F7')

    COLUNAS = [
        ('razao_social',  'Razão Social *',    35),
        ('nome_fantasia', 'Nome Fantasia',      25),
        ('cnpj',          'CNPJ',               18),
        ('tipo',          'Tipo',               15),
        ('telefone',      'Telefone',           16),
        ('email',         'E-mail',             28),
        ('endereco',      'Endereço',           35),
        ('municipio',     'Município',          20),
        ('uf',            'UF',                  5),
        ('observacoes',   'Observações',        35),
        ('ativo',         'Ativo (SIM/NAO)',    16),
    ]

    for col_idx, (_, label, width) in enumerate(COLUNAS, start=1):
        cel = ws.cell(row=1, column=col_idx, value=label)
        cel.fill = hdr_fill; cel.font = hdr_font
        cel.alignment = hdr_align; cel.border = borda
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 28

    for row_idx, f in enumerate(fornecs, start=2):
        valores = [
            f.razao_social, f.nome_fantasia or '', f.cnpj or '',
            f.tipo or '', f.telefone or '', f.email or '',
            f.endereco or '', f.municipio or '', f.uf or 'PA',
            f.observacoes or '', 'SIM' if f.ativo else 'NAO',
        ]
        row_fill = alt_fill if row_idx % 2 == 0 else None
        for col_idx, val in enumerate(valores, start=1):
            cel = ws.cell(row=row_idx, column=col_idx, value=val)
            cel.border = borda
            cel.alignment = Alignment(vertical='center')
            if row_fill: cel.fill = row_fill

    ws.freeze_panes = 'A2'
    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)

    from datetime import date as _date
    return send_file(buf, as_attachment=True,
                     download_name=f'fornecedores_{_date.today().isoformat()}.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# ── MODELO (TEMPLATE VAZIO) ──────────────────────────────
@bp.route('/fornecedores/modelo')
@login_required
def fornecedores_modelo():
    _requer_admin()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Fornecedores'

    hdr_fill  = PatternFill('solid', fgColor='F7931E')
    hdr_font  = Font(bold=True, color='FFFFFF', size=11)
    borda     = Border(
        left=Side(style='thin', color='E5E5EA'), right=Side(style='thin', color='E5E5EA'),
        top=Side(style='thin', color='E5E5EA'),  bottom=Side(style='thin', color='E5E5EA'),
    )
    COLUNAS = [
        ('Razão Social *', 35), ('Nome Fantasia', 25), ('CNPJ', 18),
        ('Tipo', 15), ('Telefone', 16), ('E-mail', 28), ('Endereço', 35),
        ('Município', 20), ('UF', 5), ('Observações', 35), ('Ativo (SIM/NAO)', 16),
    ]
    for col_idx, (label, width) in enumerate(COLUNAS, start=1):
        cel = ws.cell(row=1, column=col_idx, value=label)
        cel.fill = hdr_fill; cel.font = hdr_font
        cel.alignment = Alignment(horizontal='center', vertical='center')
        cel.border = borda
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 28

    # Linha de exemplo
    exemplo_fill = PatternFill('solid', fgColor='FFF3CD')
    for col_idx, val in enumerate(
        ['Empresa Exemplo Ltda', 'Exemplo', '00.000.000/0001-00',
         'Fornecedor', '(93) 9 9999-9999', 'contato@empresa.com.br',
         'Rua das Flores, 100', 'Santarém', 'PA', '', 'SIM'], start=1
    ):
        cel = ws.cell(row=2, column=col_idx, value=val)
        cel.fill = exemplo_fill; cel.border = borda
        cel.font = Font(italic=True, color='636366')

    # Aba de instruções
    wi = wb.create_sheet('Instruções')
    instrucoes = [
        ('INSTRUÇÕES DE PREENCHIMENTO', True),
        ('', False),
        ('• Razão Social é obrigatória (marcada com *)', False),
        ('• CNPJ: usado para identificar duplicatas — se existir, dados serão ATUALIZADOS', False),
        ('• Sem CNPJ: fornecedor sempre criado como novo', False),
        ('• UF: sigla do estado (ex: PA, AM, SP). Padrão: PA', False),
        ('• Tipo: Fornecedor | Prestador | Transportadora | Locadora | Seguradora | Outros', False),
        ('• Ativo: preencher com SIM ou NAO', False),
        ('• A linha 2 da aba Fornecedores é um exemplo — pode apagar antes de importar', False),
        ('• Não altere os cabeçalhos (linha 1)', False),
    ]
    for row_idx, (texto, negrito) in enumerate(instrucoes, start=1):
        cel = wi.cell(row=row_idx, column=1, value=texto)
        cel.font = Font(bold=negrito, size=12 if negrito else 11,
                        color='F7931E' if negrito else '000000')
    wi.column_dimensions['A'].width = 70

    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='modelo_fornecedores.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# ── IMPORT FORNECEDORES ──────────────────────────────────
@bp.route('/fornecedores/importar', methods=['POST'])
@login_required
def fornecedores_importar():
    _requer_admin()
    import re as _re

    arq = request.files.get('arquivo')
    if not arq or not arq.filename.lower().endswith(('.xlsx', '.xlsm')):
        flash('Envie um arquivo .xlsx válido.', 'danger')
        return redirect(url_for('admin.fornecedores'))

    try:
        buf = io.BytesIO(arq.read())
        wb  = openpyxl.load_workbook(buf, read_only=True, data_only=True)
        ws  = wb['Fornecedores'] if 'Fornecedores' in wb.sheetnames else wb.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
    except Exception as e:
        flash(f'Erro ao ler o arquivo: {e}', 'danger')
        return redirect(url_for('admin.fornecedores'))

    criados = atualizados = ignorados = 0
    erros = []

    for linha_num, row in enumerate(rows, start=2):
        if not any(row):
            continue

        razao = str(row[0] or '').strip()
        if not razao:
            erros.append(f'Linha {linha_num}: Razão Social vazia — ignorada.')
            ignorados += 1
            continue

        # CNPJ: normaliza para XX.XXX.XXX/XXXX-XX
        cnpj_raw  = str(row[2] or '').strip() if len(row) > 2 and row[2] else ''
        cnpj_dig  = _re.sub(r'\D', '', cnpj_raw)
        cnpj      = (f'{cnpj_dig[:2]}.{cnpj_dig[2:5]}.{cnpj_dig[5:8]}'
                     f'/{cnpj_dig[8:12]}-{cnpj_dig[12:14]}') if len(cnpj_dig) == 14 else (cnpj_raw or None)

        uf_val    = str(row[8] or 'PA').strip().upper()[:2] if len(row) > 8 and row[8] else 'PA'
        ativo_val = str(row[10] or 'SIM').strip().upper() != 'NAO' if len(row) > 10 else True

        campos = dict(
            razao_social  = razao,
            nome_fantasia = str(row[1] or '').strip() or None if len(row) > 1 else None,
            tipo          = str(row[3] or '').strip() or None if len(row) > 3 else None,
            telefone      = str(row[4] or '').strip() or None if len(row) > 4 else None,
            email         = str(row[5] or '').strip() or None if len(row) > 5 else None,
            endereco      = str(row[6] or '').strip() or None if len(row) > 6 else None,
            municipio     = str(row[7] or '').strip() or None if len(row) > 7 else None,
            uf            = uf_val,
            observacoes   = str(row[9] or '').strip() or None if len(row) > 9 else None,
            ativo         = ativo_val,
        )

        try:
            existente = Fornecedor.query.filter_by(cnpj=cnpj).first() if cnpj else None
            if existente:
                for k, v in campos.items():
                    setattr(existente, k, v)
                atualizados += 1
            else:
                db.session.add(Fornecedor(cnpj=cnpj, **campos))
                criados += 1
        except Exception as e:
            db.session.rollback()
            erros.append(f'Linha {linha_num} ({razao[:30]}): {e}')
            ignorados += 1

    try:
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        flash(f'Erro ao salvar: {e}', 'danger')
        return redirect(url_for('admin.fornecedores'))

    partes = []
    if criados:      partes.append(f'{criados} criado(s)')
    if atualizados:  partes.append(f'{atualizados} atualizado(s)')
    if ignorados:    partes.append(f'{ignorados} ignorado(s)')
    flash(f'Importação concluída: {", ".join(partes)}.', 'success' if not erros else 'warning')
    for msg in erros[:8]:
        flash(msg, 'warning')

    return redirect(url_for('admin.fornecedores'))


# ── CATÁLOGO UNIFICADO ──────────────────────────────────
@bp.route("/catalogo")
@login_required
def catalogo():
    _requer_admin()
    svcs = CatalogoServico.query.join(CategoriaServico).order_by(
        CategoriaServico.nome, CatalogoServico.nome).all()
    mats = CatalogoMaterial.query.join(CategoriaMaterial).order_by(
        CategoriaMaterial.nome, CatalogoMaterial.nome).all()
    imps = CatalogoImplemento.query.join(CategoriaImplemento).order_by(
        CategoriaImplemento.nome, CatalogoImplemento.nome).all()
    return render_template("admin/catalogo.html", servicos=svcs, peças=mats, implementos=imps)


# ── EXPORTAR CATÁLOGO EXCEL ──────────────────────────────
@bp.route('/catalogo/exportar')
@login_required
def catalogo_exportar():
    _requer_admin()
    svcs = CatalogoServico.query.join(CategoriaServico).order_by(CategoriaServico.nome, CatalogoServico.nome).all()
    mats = CatalogoMaterial.query.join(CategoriaMaterial).order_by(CategoriaMaterial.nome, CatalogoMaterial.nome).all()
    imps = CatalogoImplemento.query.join(CategoriaImplemento).order_by(CategoriaImplemento.nome, CatalogoImplemento.nome).all()

    wb = openpyxl.Workbook()

    HDR_FILL  = PatternFill("solid", fgColor="1F2937")
    HDR_FONT  = Font(bold=True, color="FFFFFF", size=11)
    HDR_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
    BORDER    = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )
    ALT_FILL = PatternFill("solid", fgColor="F9FAFB")
    SIM_FILL = PatternFill("solid", fgColor="D1FAE5")
    NAO_FILL = PatternFill("solid", fgColor="FEE2E2")

    def _estilizar_aba(ws, headers, rows):
        ws.row_dimensions[1].height = 32
        for col, h in enumerate(headers, 1):
            c = ws.cell(row=1, column=col, value=h)
            c.font = HDR_FONT; c.fill = HDR_FILL
            c.alignment = HDR_ALIGN; c.border = BORDER
        for r_idx, row in enumerate(rows, 2):
            fill = ALT_FILL if r_idx % 2 == 0 else None
            for c_idx, val in enumerate(row, 1):
                c = ws.cell(row=r_idx, column=c_idx, value=val)
                c.border = BORDER
                c.alignment = Alignment(vertical="center")
                if c_idx == len(headers):
                    if val == "Sim":
                        c.fill = SIM_FILL; c.font = Font(color="065F46", bold=True)
                    else:
                        c.fill = NAO_FILL; c.font = Font(color="991B1B", bold=True)
                elif fill:
                    c.fill = fill
        for col in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=0)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 50)

    ws_s = wb.active
    ws_s.title = "Serviços"
    _estilizar_aba(ws_s,
        ["Categoria", "Nome", "Descrição", "Tempo Est. (h)", "Valor Ref. (R$)", "Ativo"],
        [(s.categoria.nome, s.nome, s.descricao or "", float(s.tempo_estimado_h) if s.tempo_estimado_h else "",
          float(s.valor_referencia) if s.valor_referencia else "", "Sim" if s.ativo else "Não") for s in svcs])

    ws_m = wb.create_sheet("Peças")
    _estilizar_aba(ws_m,
        ["Categoria", "Nome", "Código", "Unidade", "Descrição", "Valor Ref. (R$)", "Ativo"],
        [(m.categoria.nome, m.nome, m.codigo or "", m.unidade, m.descricao or "",
          float(m.valor_referencia) if m.valor_referencia else "", "Sim" if m.ativo else "Não") for m in mats])

    ws_i = wb.create_sheet("Implementos")
    _estilizar_aba(ws_i,
        ["Categoria", "Nome", "Código", "Fabricante", "Descrição", "Valor Ref. (R$)", "Ativo"],
        [(i.categoria.nome, i.nome, i.codigo or "", i.fabricante or "", i.descricao or "",
          float(i.valor_referencia) if i.valor_referencia else "", "Sim" if i.ativo else "Não") for i in imps])

    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    return send_file(buf, as_attachment=True,
                     download_name="catalogo_frotas.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ── IMPORTAR CATÁLOGO EXCEL ──────────────────────────────
@bp.route("/catalogo/importar", methods=["POST"])
@login_required
def catalogo_importar():
    _requer_admin()
    arq = request.files.get("arquivo")
    if not arq or not arq.filename.endswith(".xlsx"):
        flash("Envie um arquivo .xlsx válido.", "danger")
        return redirect(url_for("admin.catalogo"))

    def _val(v): return str(v).strip() if v is not None else ""
    def _float(v):
        try: return float(v) if v not in (None, "") else None
        except: return None
    def _ativo(v):
        if v is None: return True
        return str(v).strip().lower() not in ("não","nao","0","false","no")

    def _ensure_cats(Model, nomes):
        """Cria categorias novas em um único commit isolado e retorna mapa nome_lower -> obj."""
        existentes = {c.nome.lower(): c for c in Model.query.all()}
        novas = []
        for nome in nomes:
            if nome.lower() not in existentes:
                obj = Model(nome=nome)
                db.session.add(obj)
                novas.append((nome.lower(), obj))
        if novas:
            db.session.commit()   # commit isolado só para categorias novas
            for key, obj in novas:
                existentes[key] = obj
        return existentes

    try:
        data = arq.read()
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        erros, criados, atualizados = [], 0, 0

        # ── SERVIÇOS ──
        if "Serviços" in wb.sheetnames:
            rows_s = [r for r in wb["Serviços"].iter_rows(min_row=2, values_only=True) if any(r)]
            cat_nomes = {_val(r[0]) for r in rows_s if _val(r[0])}
            cats = _ensure_cats(CategoriaServico, cat_nomes)
            existentes = {(s.nome.lower(), s.categoria_id): s for s in CatalogoServico.query.all()}
            novos = []
            for idx, row in enumerate(rows_s, 2):
                cat_nome, nome, descricao, tempo_h, valor, ativo_str = (list(row)+[None]*6)[:6]
                cat_nome = _val(cat_nome); nome = _val(nome)
                if not cat_nome or not nome:
                    erros.append(f"Serviços L{idx}: Categoria e Nome obrigatórios."); continue
                cat = cats.get(cat_nome.lower())
                if not cat: continue
                key = (nome.lower(), cat.id)
                if key in existentes:
                    s = existentes[key]
                    s.descricao=_val(descricao); s.tempo_estimado_h=_float(tempo_h)
                    s.valor_referencia=_float(valor); s.ativo=_ativo(ativo_str)
                    atualizados += 1
                else:
                    novos.append(CatalogoServico(categoria_id=cat.id, nome=nome,
                        descricao=_val(descricao), tempo_estimado_h=_float(tempo_h),
                        valor_referencia=_float(valor), ativo=_ativo(ativo_str)))
                    criados += 1
            db.session.add_all(novos)
            db.session.commit()

        # ── PEÇAS ──
        if "Peças" in wb.sheetnames:
            rows_m = [r for r in wb["Peças"].iter_rows(min_row=2, values_only=True) if any(r)]
            cat_nomes = {_val(r[0]) for r in rows_m if _val(r[0])}
            cats = _ensure_cats(CategoriaMaterial, cat_nomes)
            existentes = {(m.nome.lower(), m.categoria_id): m for m in CatalogoMaterial.query.all()}
            novos = []
            for idx, row in enumerate(rows_m, 2):
                cat_nome, nome, codigo, unidade, descricao, valor, ativo_str = (list(row)+[None]*7)[:7]
                cat_nome = _val(cat_nome); nome = _val(nome); unidade = _val(unidade) or "un"
                if not cat_nome or not nome:
                    erros.append(f"Peças L{idx}: Categoria e Nome obrigatórios."); continue
                cat = cats.get(cat_nome.lower())
                if not cat: continue
                key = (nome.lower(), cat.id)
                if key in existentes:
                    m = existentes[key]
                    m.codigo=_val(codigo) or None; m.unidade=unidade
                    m.descricao=_val(descricao); m.valor_referencia=_float(valor)
                    m.ativo=_ativo(ativo_str)
                    atualizados += 1
                else:
                    novos.append(CatalogoMaterial(categoria_id=cat.id, nome=nome,
                        codigo=_val(codigo) or None, unidade=unidade,
                        descricao=_val(descricao), valor_referencia=_float(valor),
                        ativo=_ativo(ativo_str)))
                    criados += 1
            db.session.add_all(novos)
            db.session.commit()

        # ── IMPLEMENTOS ──
        if "Implementos" in wb.sheetnames:
            rows_i = [r for r in wb["Implementos"].iter_rows(min_row=2, values_only=True) if any(r)]
            cat_nomes = {_val(r[0]) for r in rows_i if _val(r[0])}
            cats = _ensure_cats(CategoriaImplemento, cat_nomes)
            existentes = {(i.nome.lower(), i.categoria_id): i for i in CatalogoImplemento.query.all()}
            novos = []
            for idx, row in enumerate(rows_i, 2):
                cat_nome, nome, codigo, fabricante, descricao, valor, ativo_str = (list(row)+[None]*7)[:7]
                cat_nome = _val(cat_nome); nome = _val(nome)
                if not cat_nome or not nome:
                    erros.append(f"Implementos L{idx}: Categoria e Nome obrigatórios."); continue
                cat = cats.get(cat_nome.lower())
                if not cat: continue
                key = (nome.lower(), cat.id)
                if key in existentes:
                    i = existentes[key]
                    i.codigo=_val(codigo) or None; i.fabricante=_val(fabricante) or None
                    i.descricao=_val(descricao); i.valor_referencia=_float(valor)
                    i.ativo=_ativo(ativo_str)
                    atualizados += 1
                else:
                    novos.append(CatalogoImplemento(categoria_id=cat.id, nome=nome,
                        codigo=_val(codigo) or None, fabricante=_val(fabricante) or None,
                        descricao=_val(descricao), valor_referencia=_float(valor),
                        ativo=_ativo(ativo_str)))
                    criados += 1
            db.session.add_all(novos)
            db.session.commit()

        msg = f"Importação concluída: {criados} criados, {atualizados} atualizados."
        if erros:
            msg += f" {len(erros)} erro(s): " + " | ".join(erros[:5])
            flash(msg, "warning")
        else:
            flash(msg, "success")
    except Exception as e:
        db.session.rollback()
        flash(f"Erro ao importar: {e}", "danger")

    return redirect(url_for("admin.catalogo"))


# ── CATÁLOGO DE SERVIÇOS ─────────────────────────────────
@bp.route('/servicos')
@login_required
def servicos():
    _requer_admin()
    svcs = CatalogoServico.query.join(CategoriaServico).order_by(
        CategoriaServico.nome, CatalogoServico.nome).all()
    return render_template('admin/servicos.html', servicos=svcs)


@bp.route('/servicos/novo', methods=['GET', 'POST'])
@bp.route('/servicos/<int:id>/editar', methods=['GET', 'POST'])
@login_required
def servico_form(id=None):
    _requer_admin()
    svc  = CatalogoServico.query.get_or_404(id) if id else None
    cats = CategoriaServico.query.filter_by(ativo=True).order_by(CategoriaServico.nome).all()
    if request.method == 'POST':
        f = request.form
        try:
            if svc:
                svc.categoria_id      = int(f['categoria_id'])
                svc.nome              = f['nome']
                svc.descricao         = f.get('descricao')
                svc.tempo_estimado_h  = float(f['tempo_estimado_h']) if f.get('tempo_estimado_h') else None
                svc.valor_referencia  = float(f['valor_referencia']) if f.get('valor_referencia') else None
                svc.ativo             = 'ativo' in f
            else:
                svc = CatalogoServico(
                    categoria_id=int(f['categoria_id']),
                    nome=f['nome'],
                    descricao=f.get('descricao'),
                    tempo_estimado_h=float(f['tempo_estimado_h']) if f.get('tempo_estimado_h') else None,
                    valor_referencia=float(f['valor_referencia']) if f.get('valor_referencia') else None,
                )
                db.session.add(svc)
            db.session.commit()
            flash('Serviço salvo.', 'success')
            return redirect(url_for('admin.servicos'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro: {e}', 'danger')
    return render_template('admin/servico_form.html', svc=svc, categorias=cats)


# ── CATÁLOGO DE MATERIAIS ────────────────────────────────
@bp.route('/materiais')
@login_required
def materiais():
    _requer_admin()
    mats = CatalogoMaterial.query.join(CategoriaMaterial).order_by(
        CategoriaMaterial.nome, CatalogoMaterial.nome).all()
    return render_template('admin/materiais.html', materiais=mats)


@bp.route('/materiais/novo', methods=['GET', 'POST'])
@bp.route('/materiais/<int:id>/editar', methods=['GET', 'POST'])
@login_required
def material_form(id=None):
    _requer_admin()
    mat  = CatalogoMaterial.query.get_or_404(id) if id else None
    cats = CategoriaMaterial.query.filter_by(ativo=True).order_by(CategoriaMaterial.nome).all()
    if request.method == 'POST':
        f = request.form
        try:
            if mat:
                mat.categoria_id    = int(f['categoria_id'])
                mat.codigo          = f.get('codigo') or None
                mat.nome            = f['nome']
                mat.descricao       = f.get('descricao')
                mat.unidade         = f['unidade']
                mat.valor_referencia= float(f['valor_referencia']) if f.get('valor_referencia') else None
                mat.ativo           = 'ativo' in f
            else:
                mat = CatalogoMaterial(
                    categoria_id=int(f['categoria_id']),
                    codigo=f.get('codigo') or None,
                    nome=f['nome'],
                    descricao=f.get('descricao'),
                    unidade=f['unidade'],
                    valor_referencia=float(f['valor_referencia']) if f.get('valor_referencia') else None,
                )
                db.session.add(mat)
            db.session.commit()
            flash('Peça salva.', 'success')
            return redirect(url_for('admin.catalogo') + '#pecas')
        except Exception as e:
            db.session.rollback()
            flash(f'Erro: {e}', 'danger')
    return render_template('admin/material_form.html', mat=mat, categorias=cats)


# ── CATÁLOGO DE IMPLEMENTOS ─────────────────────────────
@bp.route('/implementos')
@login_required
def implementos():
    _requer_admin()
    imps = CatalogoImplemento.query.join(CategoriaImplemento).order_by(
        CategoriaImplemento.nome, CatalogoImplemento.nome).all()
    return render_template('admin/implementos.html', implementos=imps)


@bp.route('/implementos/novo', methods=['GET', 'POST'])
@bp.route('/implementos/<int:id>/editar', methods=['GET', 'POST'])
@login_required
def implemento_form(id=None):
    _requer_admin()
    imp  = CatalogoImplemento.query.get_or_404(id) if id else None
    cats = CategoriaImplemento.query.filter_by(ativo=True).order_by(CategoriaImplemento.nome).all()
    if request.method == 'POST':
        f = request.form
        try:
            if imp:
                imp.categoria_id     = int(f['categoria_id'])
                imp.codigo           = f.get('codigo') or None
                imp.nome             = f['nome']
                imp.descricao        = f.get('descricao')
                imp.fabricante       = f.get('fabricante') or None
                imp.valor_referencia = float(f['valor_referencia']) if f.get('valor_referencia') else None
                imp.ativo            = 'ativo' in f
            else:
                imp = CatalogoImplemento(
                    categoria_id=int(f['categoria_id']),
                    codigo=f.get('codigo') or None,
                    nome=f['nome'],
                    descricao=f.get('descricao'),
                    fabricante=f.get('fabricante') or None,
                    valor_referencia=float(f['valor_referencia']) if f.get('valor_referencia') else None,
                )
                db.session.add(imp)
            db.session.commit()
            flash('Implemento salvo.', 'success')
            return redirect(url_for('admin.implementos'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro: {e}', 'danger')
    return render_template('admin/implemento_form.html', imp=imp, categorias=cats)


# ── CHECKLIST CONFIG ─────────────────────────────────────
@bp.route('/checklist-config')
@login_required
def checklist_config():
    _requer_admin()
    itens = ChecklistItemConfig.query.order_by(ChecklistItemConfig.ordem).all()
    return render_template('admin/checklist_config.html', itens=itens)


@bp.route('/checklist-config/novo', methods=['POST'])
@login_required
def checklist_item_novo():
    _requer_admin()
    f = request.form
    try:
        item = ChecklistItemConfig(
            nome=f['nome'],
            obrigatorio='obrigatorio' in f,
            ordem=int(f.get('ordem') or 99),
        )
        db.session.add(item)
        db.session.commit()
        flash('Item adicionado.', 'success')
    except Exception as e:
        db.session.rollback()
        flash(f'Erro: {e}', 'danger')
    return redirect(url_for('admin.checklist_config'))


@bp.route('/checklist-config/<int:id>/toggle', methods=['POST'])
@login_required
def checklist_item_toggle(id):
    _requer_admin()
    item = ChecklistItemConfig.query.get_or_404(id)
    item.ativo = not item.ativo
    db.session.commit()
    return redirect(url_for('admin.checklist_config'))
