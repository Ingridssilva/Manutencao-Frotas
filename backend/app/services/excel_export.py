"""
Serviço de exportação Excel — Frota Empresa Exemplo
Gera planilhas formatadas com openpyxl para todos os módulos.
"""
from io import BytesIO
from datetime import date, datetime
from openpyxl import Workbook
from openpyxl.styles import (
    Font, Fill, PatternFill, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter
from flask import Response

# Cores da marca
COR_HEADER     = 'FF1C1C1E'   # grafite escuro
COR_SUBHEADER  = 'FFF7931E'   # laranja
COR_LINHA_PAR  = 'FFF2F2F7'   # cinza claro
COR_BRANCO     = 'FFFFFFFF'
COR_FONTE_HDR  = 'FFFFFFFF'
COR_FONTE_SUB  = 'FF1C1C1E'


def _estilo_header(ws, row: int, col_ini: int, col_fim: int, titulo: str, subtitulo: str = ''):
    """Adiciona bloco de título com identidade visual."""
    ws.row_dimensions[row].height = 32
    c = ws.cell(row=row, column=col_ini, value=titulo)
    c.font      = Font(name='Calibri', bold=True, size=16, color=COR_FONTE_HDR)
    c.fill      = PatternFill('solid', fgColor=COR_HEADER)
    c.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws.merge_cells(start_row=row, start_column=col_ini,
                   end_row=row, end_column=col_fim)
    for col in range(col_ini + 1, col_fim + 1):
        ws.cell(row=row, column=col).fill = PatternFill('solid', fgColor=COR_HEADER)

    if subtitulo:
        ws.row_dimensions[row + 1].height = 20
        c2 = ws.cell(row=row + 1, column=col_ini, value=subtitulo)
        c2.font      = Font(name='Calibri', size=9, color='FF8E8E93')
        c2.fill      = PatternFill('solid', fgColor=COR_HEADER)
        c2.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        ws.merge_cells(start_row=row + 1, start_column=col_ini,
                       end_row=row + 1, end_column=col_fim)
        for col in range(col_ini + 1, col_fim + 1):
            ws.cell(row=row + 1, column=col).fill = PatternFill('solid', fgColor=COR_HEADER)
        return row + 2
    return row + 1


def _estilo_colunas(ws, headers: list, row: int):
    """Estiliza linha de cabeçalho das colunas."""
    ws.row_dimensions[row].height = 22
    borda = Border(
        bottom=Side(style='thin', color='FFF7931E'),
    )
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font      = Font(name='Calibri', bold=True, size=10, color=COR_FONTE_HDR)
        c.fill      = PatternFill('solid', fgColor='FF2C2C2E')
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border    = borda


def _linha(ws, row: int, valores: list, par: bool = False):
    """Preenche uma linha de dados."""
    fill = PatternFill('solid', fgColor=COR_LINHA_PAR if par else COR_BRANCO)
    for i, v in enumerate(valores, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font      = Font(name='Calibri', size=10)
        c.fill      = fill
        c.alignment = Alignment(vertical='center', wrap_text=False)


def _autofit(ws, min_width=10, max_width=50):
    """Ajusta largura de colunas automaticamente."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_width, max(min_width, max_len + 2))


def _response(wb: Workbook, filename: str) -> Response:
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(
        buf.read(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename="{filename}"'}
    )


def _brl(v) -> str:
    try:
        return f'R$ {float(v):,.2f}'.replace(',','X').replace('.', ',').replace('X','.')
    except Exception:
        return '—'


def _fmt_date(d) -> str:
    if not d:
        return '—'
    try:
        return d.strftime('%d/%m/%Y')
    except Exception:
        return str(d)


# ──────────────────────────────────────────────────────────
#  EXPORTAÇÕES POR MÓDULO
# ──────────────────────────────────────────────────────────

def exportar_os(os_list) -> Response:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Ordens de Serviço'
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A4'

    headers = ['Número','Data Abertura','Veículo','Placa','Tipo','Manutenção',
               'Status','Fornecedor','KM Abertura','Data Conclusão','Custo Total','Observações']

    row = _estilo_header(ws, 1, 1, len(headers),
                         '🔧 Ordens de Serviço — Empresa Exemplo',
                         f'Exportado em {date.today().strftime("%d/%m/%Y")} · {len(os_list)} registros')
    _estilo_colunas(ws, headers, row)
    row += 1

    for i, o in enumerate(os_list):
        _linha(ws, row, [
            o.numero or '—',
            _fmt_date(o.data_abertura),
            (o.veiculo.descricao or o.veiculo.modelo or '') if o.veiculo else '—',
            o.veiculo.placa if o.veiculo else '—',
            o.tipo_os or '—',
            o.tipo_manutencao or '—',
            o.status or '—',
            (o.fornecedor.razao_social if o.fornecedor else '—'),
            o.km_abertura or '—',
            _fmt_date(o.data_conclusao),
            float(o.custo_total) if o.custo_total else 0,
            o.observacoes or '',
        ], par=i % 2 == 0)
        row += 1

    # Formata coluna custo como moeda
    for r in range(4, row):
        ws.cell(r, 11).number_format = 'R$ #,##0.00'

    _autofit(ws)
    return _response(wb, f'OS_{date.today().strftime("%Y%m%d")}.xlsx')


def exportar_veiculos(veiculos) -> Response:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Veículos'
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A4'

    headers = ['Placa','Descrição','Marca','Modelo','Ano','Cor','Tipo Veículo',
               'Município Base','Propriedade','KM Atual','Status','Fornecedor Locação']

    row = _estilo_header(ws, 1, 1, len(headers),
                         '🚗 Frota de Veículos — Empresa Exemplo',
                         f'Exportado em {date.today().strftime("%d/%m/%Y")} · {len(veiculos)} veículos')
    _estilo_colunas(ws, headers, row)
    row += 1

    for i, v in enumerate(veiculos):
        _linha(ws, row, [
            v.placa,
            v.descricao or '—',
            v.marca or '—',
            v.modelo or '—',
            v.ano_fabricacao or '—',
            v.cor or '—',
            v.tipo_veiculo or '—',
            v.municipio_base or '—',
            v.tipo_propriedade or '—',
            v.km_atual or 0,
            v.status or '—',
            (v.fornecedor_locacao.razao_social if v.fornecedor_locacao else '—'),
        ], par=i % 2 == 0)
        row += 1

    _autofit(ws)
    return _response(wb, f'Veiculos_{date.today().strftime("%Y%m%d")}.xlsx')


def exportar_motoristas(motoristas) -> Response:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Motoristas'
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A4'

    headers = ['Nome','CPF','Matrícula','Telefone','Município','CNH Número',
               'CNH Categoria','CNH Validade','CNH Status','Pontos CNH',
               'Veículo Principal','Ativo']

    row = _estilo_header(ws, 1, 1, len(headers),
                         '👤 Motoristas — Empresa Exemplo',
                         f'Exportado em {date.today().strftime("%d/%m/%Y")} · {len(motoristas)} motoristas')
    _estilo_colunas(ws, headers, row)
    row += 1

    for i, m in enumerate(motoristas):
        _linha(ws, row, [
            m.nome,
            m.cpf or '—',
            m.matricula or '—',
            m.telefone or '—',
            m.municipio or '—',
            m.cnh_numero or '—',
            m.cnh_categoria or '—',
            _fmt_date(m.cnh_validade),
            'VENCIDA' if m.cnh_vencida else (m.cnh_status or '—'),
            m.cnh_pontos or 0,
            (m.veiculo_principal.placa if m.veiculo_principal else '—'),
            'Sim' if m.ativo else 'Não',
        ], par=i % 2 == 0)

        # Destaca CNH vencida em vermelho
        if m.cnh_vencida:
            ws.cell(row, 8).font = Font(name='Calibri', size=10, color='FFC62828', bold=True)

        row += 1

    _autofit(ws)
    return _response(wb, f'Motoristas_{date.today().strftime("%Y%m%d")}.xlsx')


def exportar_pneus(pneus) -> Response:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Pneus'
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A4'

    headers = ['Código','Marca','Modelo','Dimensão','Tipo','DOT','KM Rodado',
               'KM Limite','Vida %','Status Desgaste','Recapagens',
               'Veículo Atual','Posição','Status']

    row = _estilo_header(ws, 1, 1, len(headers),
                         '⭕ Controle de Pneus — Empresa Exemplo',
                         f'Exportado em {date.today().strftime("%d/%m/%Y")} · {len(pneus)} pneus')
    _estilo_colunas(ws, headers, row)
    row += 1

    for i, p in enumerate(pneus):
        pos = p.posicao_atual
        _linha(ws, row, [
            p.codigo or '—',
            p.marca or '—',
            p.modelo or '—',
            p.dimensao or '—',
            p.tipo,
            p.dot or '—',
            p.km_rodado,
            p.km_limite or 0,
            p.pct_vida,
            p.status_desgaste,
            p.vezes_recapado,
            pos.veiculo.placa if pos and pos.veiculo else '—',
            pos.posicao_display if pos else '—',
            p.status,
        ], par=i % 2 == 0)

        # Cor na coluna vida %
        pct_cell = ws.cell(row, 9)
        if p.pct_vida >= 95:
            pct_cell.font = Font(name='Calibri', size=10, color='FFC62828', bold=True)
        elif p.pct_vida >= 80:
            pct_cell.font = Font(name='Calibri', size=10, color='FFE65100', bold=True)

        row += 1

    ws.column_dimensions['I'].number_format = '0"%"'
    _autofit(ws)
    return _response(wb, f'Pneus_{date.today().strftime("%Y%m%d")}.xlsx')


def exportar_fluidos(trocas) -> Response:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Fluidos'
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A4'

    headers = ['Veículo','Placa','Tipo Fluido','Data Troca','KM Troca',
               'KM Próxima Troca','KM Restante','Próxima Data','Marca/Produto',
               'Viscosidade','Qtd (L)','Custo','Status']

    row = _estilo_header(ws, 1, 1, len(headers),
                         '🛢 Controle de Fluidos — Empresa Exemplo',
                         f'Exportado em {date.today().strftime("%d/%m/%Y")} · {len(trocas)} registros')
    _estilo_colunas(ws, headers, row)
    row += 1

    for i, t in enumerate(trocas):
        kr = t.km_restante
        _linha(ws, row, [
            (t.veiculo.descricao or t.veiculo.modelo or '') if t.veiculo else '—',
            t.veiculo.placa if t.veiculo else '—',
            t.tipo_display,
            _fmt_date(t.data_troca),
            t.km_troca or 0,
            t.km_proxima_troca or 0,
            kr if kr is not None else '—',
            _fmt_date(t.proxima_data),
            t.marca_produto or '—',
            t.viscosidade or '—',
            float(t.quantidade_litros) if t.quantidade_litros else 0,
            float(t.custo) if t.custo else 0,
            t.status_alerta,
        ], par=i % 2 == 0)

        ws.cell(row, 12).number_format = 'R$ #,##0.00'
        st = t.status_alerta
        if st in ('vencido', 'critico'):
            ws.cell(row, 13).font = Font(name='Calibri', size=10, color='FFC62828', bold=True)
        elif st == 'atencao':
            ws.cell(row, 13).font = Font(name='Calibri', size=10, color='FFE65100', bold=True)

        row += 1

    _autofit(ws)
    return _response(wb, f'Fluidos_{date.today().strftime("%Y%m%d")}.xlsx')


def exportar_multas(multas) -> Response:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Multas'
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A4'

    headers = ['Auto de Infração','Veículo','Placa','Motorista','Data Infração',
               'Município','Descrição','Pontos','Valor Original','Valor Desconto',
               'Valor Pago','Vencimento','Status','Responsável']

    row = _estilo_header(ws, 1, 1, len(headers),
                         '🚨 Controle de Multas — Empresa Exemplo',
                         f'Exportado em {date.today().strftime("%d/%m/%Y")} · {len(multas)} multas')
    _estilo_colunas(ws, headers, row)
    row += 1

    for i, m in enumerate(multas):
        _linha(ws, row, [
            m.auto_infracao or '—',
            (m.veiculo.descricao or m.veiculo.modelo or '') if m.veiculo else '—',
            m.veiculo.placa if m.veiculo else '—',
            m.motorista.nome if m.motorista else '—',
            _fmt_date(m.data_infracao),
            m.municipio or '—',
            (m.descricao or '')[:80],
            m.pontos or 0,
            float(m.valor_original) if m.valor_original else 0,
            float(m.valor_desconto) if m.valor_desconto else 0,
            float(m.valor_pago) if m.valor_pago else 0,
            _fmt_date(m.data_vencimento),
            m.status or '—',
            m.responsavel or '—',
        ], par=i % 2 == 0)

        for col in (9, 10, 11):
            ws.cell(row, col).number_format = 'R$ #,##0.00'

        row += 1

    _autofit(ws)
    return _response(wb, f'Multas_{date.today().strftime("%Y%m%d")}.xlsx')
